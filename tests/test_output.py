import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from tests.conftest import (
    FetcherTestCase, scholarly_mod, _DummyNav, _DummyIterator,
    patch, StringIO, json, tempfile,
)
import types
import unittest
import shutil
from code import scholar_citation
from code.output_state import PaperFetchState
class PapersRootTests(unittest.TestCase):
    """论文数据树的物理位置（2026-10-15 起在阶段 1 的 output 下）。

    `papers_root()` / `ensure_papers_root()` 是**唯一**给出这个路径的地方；消费方
    （阶段 2 / 3.x / 阶段 4）一律经各自的 `papers` 符号链接，不直接拼这个路径。
    """

    def test_path_shape(self):
        from code.config import papers_root
        self.assertEqual(papers_root("/tmp/out", "HF3-E9kAAAAJ"),
                         "/tmp/out/author_HF3-E9kAAAAJ_papers")

    def test_created_and_idempotent_without_touching_content(self):
        from code.config import ensure_papers_root, papers_root
        with tempfile.TemporaryDirectory() as tmp:
            path = ensure_papers_root(tmp, "A1")
            self.assertTrue(os.path.isdir(path))
            self.assertEqual(path, papers_root(tmp, "A1"))
            marker = os.path.join(path, "001-Some_Paper")
            os.makedirs(marker)
            with open(os.path.join(marker, "keep.txt"), "w", encoding="utf-8") as handle:
                handle.write("downstream data")
            # 再跑一次：路径不变、目录里的东西一个都不能动（阶段 2/3.x 往里写过东西）
            self.assertEqual(ensure_papers_root(tmp, "A1"), path)
            self.assertTrue(os.path.isfile(os.path.join(marker, "keep.txt")))

    def test_main_creates_it(self):
        """`main()` 真的接了线（不是只有 helper）。"""
        from code import config as config_module
        source = open(os.path.join(os.path.dirname(os.path.dirname(
            os.path.abspath(__file__))), "code", "scholar_citation.py"),
            encoding="utf-8").read()
        self.assertIn("ensure_papers_root(args.output_dir, author_id)", source)
        self.assertIs(config_module.ensure_papers_root, config_module.ensure_papers_root)


class OutputAndReconciliationTests(FetcherTestCase):
    def test_save_output_writes_excel_run_metadata_from_json_payload(self):
        pub = {
            "no": 1,
            "title": "Paper One",
            "num_citations": 1,
            "year": "2024",
            "venue": "Venue",
        }
        result = {
            "pub": pub,
            "citations": [
                {"title": "Citing Paper", "authors": "A", "venue": "CV", "year": "2024", "url": "https://example.com/cite"}
            ],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            self.fetcher.profile_json = os.path.join(tmpdir, "author_test-author_profile.json")
            self.fetcher.out_json = os.path.join(tmpdir, "author_test-author_paper_citations.json")
            self.fetcher.out_xlsx = os.path.join(tmpdir, "author_test-author_paper_citations.xlsx")
            self.fetcher._run_start_time = 0
            with open(self.fetcher.profile_json, "w", encoding="utf-8") as f:
                json.dump({"publications": [pub]}, f)

            captured = {}
            original_workbook = scholar_citation.openpyxl.Workbook

            class CapturingWorkbook(original_workbook):
                def __init__(self):
                    super().__init__()
                    captured["workbook"] = self

            with patch.object(scholar_citation.openpyxl, "Workbook", CapturingWorkbook), \
                 patch.object(scholar_citation, "datetime") as fake_datetime:
                fake_datetime.now.return_value = types.SimpleNamespace(
                    isoformat=lambda: "2026-04-07T12:34:56",
                    strftime=lambda fmt: "20260407_123456",
                    year=2026,
                )
                self.fetcher._save_output([result])

            with open(self.fetcher.out_json, "r", encoding="utf-8") as f:
                payload = json.load(f)

        workbook = captured["workbook"]
        self.assertEqual(payload["author_id"], "test-author")
        self.assertEqual(payload["fetch_time"], "2026-04-07T12:34:56")
        self.assertEqual(payload["total_papers"], 1)
        self.assertEqual(payload["total_citations_collected"], 1)
        self.assertEqual(len(workbook.sheets), 3)
        self.assertEqual(workbook.sheets[2].title, "Run Metadata")
        metadata_sheet = workbook.sheets[2]
        self.assertEqual(metadata_sheet.cells[(1, 1)].value, "Author ID")
        self.assertEqual(metadata_sheet.cells[(1, 2)].value, payload["author_id"])
        self.assertEqual(metadata_sheet.cells[(2, 1)].value, "Fetch Time")
        self.assertEqual(metadata_sheet.cells[(2, 2)].value, payload["fetch_time"])
        self.assertEqual(metadata_sheet.cells[(3, 1)].value, "Total Papers")
        self.assertEqual(metadata_sheet.cells[(3, 2)].value, payload["total_papers"])
        self.assertEqual(metadata_sheet.cells[(4, 1)].value, "Total Citations Collected")
        self.assertEqual(metadata_sheet.cells[(4, 2)].value, payload["total_citations_collected"])

        citations_sheet = workbook.sheets[1]
        self.assertEqual(citations_sheet.title, "All Citations")
        self.assertEqual(citations_sheet.cells[(2, 5)].value, payload["papers"][0]["citations"][0]["year"])

    def test_save_output_includes_fetch_state_from_output_state(self):
        pub = {
            "no": 1,
            "title": "Paper With State",
            "num_citations": 150,
            "year": "2020",
            "venue": "Venue",
        }
        result = {
            "pub": pub,
            "citations": [
                {"title": "Citing Paper", "authors": "A", "venue": "CV", "year": "2024", "url": "https://example.com/cite"}
            ],
        }
        state_dict = {
            "title": "Paper With State",
            "num_citations_on_scholar": 150,
            "complete_fetch_attempt": True,
            "fetch_strategy": "year",
            "year_fetch_diagnostics": {
                "2024": {
                    "year": 2024,
                    "histogram_count": 50, "cached_total": 1,
                    "seen_total": 50, "dedup_count": 0,
                    "termination_reason": "short_page_stop",
                },
                "histogram_total": 50,
                "scholar_total": 150,
                "cached_total": 1,
                "cached_year_total": 1,
                "seen_total": 50,
                "dedup_count": 0,
            },
            "citations": result["citations"],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            self.fetcher.profile_json = os.path.join(tmpdir, "author_test-author_profile.json")
            self.fetcher.out_json = os.path.join(tmpdir, "author_test-author_paper_citations.json")
            self.fetcher.out_xlsx = os.path.join(tmpdir, "author_test-author_paper_citations.xlsx")
            self.fetcher._run_start_time = 0
            with open(self.fetcher.profile_json, "w", encoding="utf-8") as f:
                json.dump({"publications": [pub]}, f)
            # Set output state (not cache)
            self.fetcher._paper_states = {
                "Paper With State": PaperFetchState.from_dict(state_dict),
            }

            with patch.object(scholar_citation.openpyxl, "Workbook", return_value=scholar_citation.openpyxl.Workbook()), \
                 patch.object(scholar_citation, "datetime") as fake_datetime:
                fake_datetime.now.return_value = types.SimpleNamespace(
                    isoformat=lambda: "2026-05-01T12:00:00",
                    strftime=lambda fmt: "20260501_120000",
                    year=2026,
                )
                self.fetcher._save_output([result])

            with open(self.fetcher.out_json, "r", encoding="utf-8") as f:
                payload = json.load(f)

        paper = payload["papers"][0]
        self.assertIn("_fetch_state", paper)
        fs = paper["_fetch_state"]
        self.assertEqual(fs["num_citations_on_scholar"], 150)
        self.assertTrue(fs["complete_fetch_attempt"])
        self.assertIn("year_fetch_diagnostics", fs)
        self.assertIn("year_records", fs)
        yr = fs["year_records"][0]
        self.assertEqual(yr["histogram_count"], 50)
        # citations should NOT be duplicated inside _fetch_state
        self.assertNotIn("citations", fs)

    # -- interrupted run: keep this run's progress -------------------------

    _INTERRUPTED_PUB = {
        "no": 1,
        "title": "Interrupted Paper",
        "num_citations": 95,
        "year": "2022",
        "venue": "V",
    }

    @staticmethod
    def _stale_complete_state():
        """What the previous run persisted: aggregate claims complete."""
        return {
            "title": "Interrupted Paper",
            "fetch_strategy": "year",
            "num_citations_on_scholar": 91,
            "complete_fetch_attempt": True,
            "year_fetch_diagnostics": {
                "2022": {"year": 2022, "histogram_count": 6, "cached_total": 7,
                         "seen_total": 7, "dedup_count": 0},
                "2026": {"year": 2026, "histogram_count": 4, "cached_total": 4,
                         "seen_total": 4, "dedup_count": 0},
                "histogram_total": 91, "seen_total": 93,
            },
        }

    @staticmethod
    def _mid_paper_snapshot(partial_citations):
        """What save_progress leaves behind when the fetch is interrupted.

        2026 was probed (5) and only 4 are cached, so the year is short — the
        aggregate (93 >= 91) still hides it.
        """
        return {
            "title": "Interrupted Paper",
            "fetch_strategy": "year",
            "num_citations_on_scholar": 95,
            "complete_fetch_attempt": False,
            "year_records": [
                {"year": 2022, "histogram_count": 6, "cached_total": 7,
                 "seen_total": 7, "dedup_count": 0},
                {"year": 2026, "histogram_count": 5, "cached_total": 4,
                 "seen_total": 4, "dedup_count": 0},
            ],
            "year_fetch_diagnostics": {"histogram_total": 91, "seen_total": 93},
            "citations": partial_citations,
        }

    def _save_and_load(self, state, with_snapshot, mid_snapshot=None, pub=None):
        """Persist one paper via _save_output and return (entry, partial_citations).

        The temporary directory survives the call so the caller can reload the
        written output (see the round-trip test below).
        """
        partial_citations = [
            {"title": f"New-2026-{idx}", "authors": "A", "venue": "V", "year": "2026", "url": f"u{idx}"}
            for idx in range(3)
        ]
        pub = dict(pub or self._INTERRUPTED_PUB)
        tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmpdir, ignore_errors=True)
        self.fetcher.profile_json = os.path.join(tmpdir, "author_test-author_profile.json")
        self.fetcher.out_json = os.path.join(tmpdir, "author_test-author_paper_citations.json")
        self.fetcher.out_xlsx = os.path.join(tmpdir, "author_test-author_paper_citations.xlsx")
        self.fetcher._run_start_time = 0
        with open(self.fetcher.profile_json, "w", encoding="utf-8") as f:
            json.dump({"publications": [pub]}, f)
        self.fetcher._paper_states = {
            "Interrupted Paper": PaperFetchState.from_dict(state),
        }
        if with_snapshot:
            if mid_snapshot is None:
                mid_snapshot = self._mid_paper_snapshot(partial_citations)
            self.fetcher._mid_paper_state = {"Interrupted Paper": mid_snapshot}
        else:
            self.fetcher._mid_paper_state = {}

        with patch("sys.stdout", new_callable=StringIO):
            self.fetcher._save_output([None])

        with open(self.fetcher.out_json, "r", encoding="utf-8") as f:
            payload = json.load(f)
        return payload["papers"][0], partial_citations

    def test_save_output_keeps_interrupted_run_progress(self):
        """The in-run snapshot must win: it holds the live probe (so 2026 stays
        visibly short) and the citations fetched before the interrupt.  Without
        it the stale start-of-run state would persist fetch_complete=True and
        the un-fetched citations would never be resumed."""
        paper, partial_citations = self._save_and_load(
            self._stale_complete_state(), with_snapshot=True)
        self.assertFalse(paper["fetch_complete"])
        self.assertIs(paper["_fetch_state"]["complete_fetch_attempt"], False)
        self.assertEqual([c["title"] for c in paper["citations"]],
                         [c["title"] for c in partial_citations])
        records = {r["year"]: r for r in paper["_fetch_state"]["year_records"]}
        self.assertEqual(records[2026]["histogram_count"], 5)
        self.assertEqual(records[2026]["seen_total"], 4)

    def test_save_output_uses_stale_state_without_in_run_snapshot(self):
        """Baseline for the test above: with no in-run snapshot the stale state
        is persisted as complete, hiding the short year."""
        paper, _ = self._save_and_load(self._stale_complete_state(), with_snapshot=False)
        self.assertTrue(paper["fetch_complete"])
        records = {r["year"]: r for r in paper["_fetch_state"]["year_records"]}
        self.assertEqual(records[2026]["histogram_count"], 4)

    def test_save_output_persists_stall_counters_alongside_in_run_snapshot(self):
        """Stall counters must survive a save that uses the in-run snapshot.

        save_progress() builds the snapshot without stall fields and the counters
        are recorded on _paper_states after the fetch, so persisting the snapshot
        verbatim reset stall_count to 0 on every run and `stalled` could never be
        reached.  The signature is taken from the snapshot-derived state (the
        live probe is what the next run compares against), so this also pins that
        it does not drift across the write/read cycle.
        """
        snapshot = self._mid_paper_snapshot([
            {"title": f"New-2026-{idx}", "authors": "A", "venue": "V",
             "year": "2026", "url": f"u{idx}"}
            for idx in range(3)
        ])
        signature = PaperFetchState.from_dict(snapshot).current_outcome_signature()
        self.assertTrue(signature)
        state = self._stale_complete_state()
        state["stall_signature"] = signature
        state["stall_count"] = 2

        paper, _ = self._save_and_load(state, with_snapshot=True, mid_snapshot=snapshot)

        saved = paper["_fetch_state"]
        self.assertEqual(saved["stall_signature"], signature)
        self.assertEqual(saved["stall_count"], 2)
        reloaded = PaperFetchState.from_dict(saved)
        self.assertTrue(reloaded.is_stalled())

    def test_save_output_keeps_json_when_excel_write_fails(self):
        """Excel is a convenience view; its failure must not abort the run."""
        pub = dict(self._INTERRUPTED_PUB)
        tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmpdir, ignore_errors=True)
        self.fetcher.profile_json = os.path.join(tmpdir, "author_test-author_profile.json")
        self.fetcher.out_json = os.path.join(tmpdir, "author_test-author_paper_citations.json")
        self.fetcher.out_xlsx = os.path.join(tmpdir, "author_test-author_paper_citations.xlsx")
        self.fetcher._run_start_time = 0
        with open(self.fetcher.profile_json, "w", encoding="utf-8") as f:
            json.dump({"publications": [pub]}, f)
        self.fetcher._paper_states = {
            "Interrupted Paper": PaperFetchState.from_dict(self._stale_complete_state()),
        }
        self.fetcher._mid_paper_state = {}

        with patch.object(self.fetcher, "_save_xlsx",
                          side_effect=RuntimeError("bad character")), \
             patch("sys.stdout", new_callable=StringIO) as stdout:
            self.fetcher._save_output([None])

        self.assertIn("could not write", stdout.getvalue())
        with open(self.fetcher.out_json, "r", encoding="utf-8") as f:
            payload = json.load(f)
        self.assertEqual(payload["total_papers"], 1)

    def test_save_output_is_idempotent_for_in_flight_progress(self):
        """Saving twice must not lose the in-flight paper's citations.

        The retry-exhausted path saves before sys.exit(1) and run()'s finally
        saves again; the first save used to clear the in-run snapshots, so the
        second wrote the stale start-of-run state without this run's citations.
        """
        state = self._stale_complete_state()
        paper, partial_citations = self._save_and_load(state, with_snapshot=True)
        self.assertEqual(len(paper["citations"]), len(partial_citations))

        with patch("sys.stdout", new_callable=StringIO):
            self.fetcher._save_output([None])

        with open(self.fetcher.out_json, "r", encoding="utf-8") as f:
            payload = json.load(f)
        self.assertEqual([c["title"] for c in payload["papers"][0]["citations"]],
                         [c["title"] for c in partial_citations])

    def test_write_json_atomic_keeps_the_previous_file_on_failure(self):
        from code.common import write_json_atomic

        with tempfile.TemporaryDirectory() as tmpdir:
            target = os.path.join(tmpdir, "state.json")
            write_json_atomic(target, {"ok": 1})
            with open(target, "r", encoding="utf-8") as f:
                self.assertEqual(json.load(f), {"ok": 1})

            class Unserializable:
                pass

            with self.assertRaises(TypeError):
                write_json_atomic(target, {"bad": Unserializable()})

            with open(target, "r", encoding="utf-8") as f:
                self.assertEqual(json.load(f), {"ok": 1})
            leftovers = [name for name in os.listdir(tmpdir) if name.endswith(".tmp")]
            self.assertEqual(leftovers, [])

    def test_corrupt_output_is_backed_up_and_treated_as_missing(self):
        from code.output_state import load_paper_states
        with tempfile.TemporaryDirectory() as tmpdir:
            target = os.path.join(tmpdir, "author_x_paper_citations.json")
            with open(target, "w", encoding="utf-8") as f:
                f.write('{"papers": [{"pub": {"title": "T"')
            with patch("sys.stdout", new_callable=StringIO) as stdout:
                states = load_paper_states(target)
            self.assertEqual(states, {})
            self.assertTrue(os.path.exists(target + ".corrupt"))
            self.assertIn("could not be read", stdout.getvalue())

    def test_skip_zero_keeps_previously_collected_citations(self):
        """A profile reporting 0 citations must not delete the stored list."""
        pub = {
            "no": 1,
            "title": "Zeroed Paper",
            "num_citations": 0,
            "year": "2024",
            "venue": "V",
        }
        kept = [
            {"title": f"Kept-{idx}", "authors": "A", "venue": "V", "year": "2024",
             "url": f"u{idx}", "cites_id": f"cid-{idx}"}
            for idx in range(3)
        ]
        tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmpdir, ignore_errors=True)
        self.fetcher.profile_json = os.path.join(tmpdir, "author_test-author_profile.json")
        self.fetcher.out_json = os.path.join(tmpdir, "author_test-author_paper_citations.json")
        self.fetcher.out_xlsx = os.path.join(tmpdir, "author_test-author_paper_citations.xlsx")
        self.fetcher._run_start_time = 0
        with open(self.fetcher.profile_json, "w", encoding="utf-8") as f:
            json.dump({"publications": [pub]}, f)
        with open(self.fetcher.out_json, "w", encoding="utf-8") as f:
            json.dump({"papers": [{
                "pub": pub,
                "citations": kept,
                "fetch_complete": True,
                "_fetch_state": {
                    "title": pub["title"], "fetch_strategy": "direct",
                    "num_citations_on_scholar": 3,
                    "complete_fetch_attempt": True,
                    "direct_fetch_diagnostics": {
                        "scholar_total": 3, "cached_total": 3, "seen_total": 3,
                        "dedup_count": 0, "termination_reason": "iterator_exhausted"},
                },
            }]}, f)
        self.fetcher._session_patched = True

        with patch("sys.stdout", new_callable=StringIO) as stdout:
            self.fetcher.run()

        self.assertIn("keeping 3 cached", stdout.getvalue())
        with open(self.fetcher.out_json, "r", encoding="utf-8") as f:
            payload = json.load(f)
        self.assertEqual([c["title"] for c in payload["papers"][0]["citations"]],
                         [c["title"] for c in kept])

    def test_save_output_deduplicates_rows_with_the_same_identity(self):
        """Scholar can list one cluster twice with contradictory metadata.

        Both rows used to be written, inflating the count; the output must keep
        one row and prefer the richer metadata.
        """
        pub = {"no": 1, "title": "Paper", "num_citations": 3,
               "year": "2024", "venue": "V"}
        citations = [
            {"title": "Dup Paper", "authors": "A", "venue": "NA", "year": "2023",
             "url": "u1", "cites_id": "cid-1"},
            {"title": "Dup Paper", "authors": "A", "venue": "Journal", "year": "2022",
             "url": "u1", "cites_id": "cid-1"},
            {"title": "Other", "authors": "B", "venue": "V2", "year": "2024",
             "url": "u2", "cites_id": "cid-2"},
        ]
        tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmpdir, ignore_errors=True)
        self.fetcher.profile_json = os.path.join(tmpdir, "author_test-author_profile.json")
        self.fetcher.out_json = os.path.join(tmpdir, "author_test-author_paper_citations.json")
        self.fetcher.out_xlsx = os.path.join(tmpdir, "author_test-author_paper_citations.xlsx")
        self.fetcher._run_start_time = 0
        with open(self.fetcher.profile_json, "w", encoding="utf-8") as f:
            json.dump({"publications": [pub]}, f)
        self.fetcher._paper_states = {}
        self.fetcher._mid_paper_state = {}
        self.fetcher._session_patched = True

        with patch("sys.stdout", new_callable=StringIO) as stdout:
            self.fetcher._save_output([{"pub": pub, "citations": citations}])

        self.assertIn("Deduplicated 1 duplicate citation row(s)", stdout.getvalue())
        with open(self.fetcher.out_json, "r", encoding="utf-8") as f:
            payload = json.load(f)
        kept = payload["papers"][0]["citations"]
        self.assertEqual(len(kept), 2)
        self.assertEqual(kept[0]["venue"], "Journal")
        self.assertEqual(kept[0]["year"], "2022")
        self.assertEqual(payload["total_citations_collected"], 2)

    def test_save_output_deduplicates_by_title_without_cites_id(self):
        pub = {"no": 1, "title": "Paper", "num_citations": 2,
               "year": "2024", "venue": "V"}
        citations = [
            {"title": "Same Title", "authors": "A", "venue": "V", "year": "N/A",
             "url": "u1", "cites_id": None},
            {"title": "Same Title", "authors": "A", "venue": "V", "year": "2020",
             "url": "u2", "cites_id": None},
        ]
        tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmpdir, ignore_errors=True)
        self.fetcher.profile_json = os.path.join(tmpdir, "author_test-author_profile.json")
        self.fetcher.out_json = os.path.join(tmpdir, "author_test-author_paper_citations.json")
        self.fetcher.out_xlsx = os.path.join(tmpdir, "author_test-author_paper_citations.xlsx")
        self.fetcher._run_start_time = 0
        with open(self.fetcher.profile_json, "w", encoding="utf-8") as f:
            json.dump({"publications": [pub]}, f)
        self.fetcher._paper_states = {}
        self.fetcher._mid_paper_state = {}
        self.fetcher._session_patched = True

        with patch("sys.stdout", new_callable=StringIO):
            self.fetcher._save_output([{"pub": pub, "citations": citations}])

        with open(self.fetcher.out_json, "r", encoding="utf-8") as f:
            payload = json.load(f)
        kept = payload["papers"][0]["citations"]
        self.assertEqual(len(kept), 1)
        self.assertEqual(kept[0]["year"], "2020")  # richer row wins

    def test_citations_xlsx_round_trips_with_real_openpyxl(self):
        """The Excel writers are only ever exercised against a stub.

        tests/conftest.py replaces openpyxl, so a broken openpyxl API call (or a
        value openpyxl rejects) would never surface in the suite.  This test
        temporarily restores the real library and reopens what it wrote.
        """
        import importlib
        stubbed = {name: sys.modules.pop(name)
                   for name in list(sys.modules)
                   if name == 'openpyxl' or name.startswith('openpyxl.')}
        try:
            real_openpyxl = importlib.import_module('openpyxl')
            from openpyxl.styles import Alignment, Font, PatternFill
            self.assertTrue(hasattr(real_openpyxl, 'load_workbook'))

            from code.citation_io import save_citations_xlsx

            pub = {"no": 1, "title": "Paper", "year": "2024", "venue": "V",
                   "num_citations": 2}
            results = [{
                "pub": pub,
                "fetch_complete": True,
                "citations": [
                    {"title": "Cited A", "authors": "A", "venue": "V1",
                     "year": "2023", "url": "https://example.com/a",
                     "cites_id": "cid-a"},
                    {"title": "Cited B", "authors": "B", "venue": "V2",
                     "year": "N/A", "url": "N/A", "cites_id": None},
                ],
            }]
            tmpdir = tempfile.mkdtemp()
            self.addCleanup(shutil.rmtree, tmpdir, ignore_errors=True)
            path = os.path.join(tmpdir, "citations.xlsx")
            save_citations_xlsx(
                path, results, "test-author",
                metadata={"author_id": "test-author", "fetch_time": "t",
                          "total_papers": 1, "total_citations_collected": 2},
                openpyxl_module=real_openpyxl,
                font_cls=Font, pattern_fill_cls=PatternFill,
                alignment_cls=Alignment,
            )
            self.assertTrue(os.path.exists(path))

            workbook = real_openpyxl.load_workbook(path)
            self.assertEqual(workbook.sheetnames,
                             ["Summary", "All Citations", "Run Metadata"])
            summary = workbook["Summary"]
            self.assertEqual(summary.cell(row=1, column=1).value, "No.")
            self.assertEqual(summary.cell(row=2, column=2).value, "Paper")
            self.assertEqual(summary.cell(row=2, column=6).value, 2)
            citations = workbook["All Citations"]
            self.assertEqual(citations.cell(row=2, column=2).value, "Cited A")
            self.assertEqual(citations.cell(row=3, column=2).value, "Cited B")
            self.assertEqual(workbook["Run Metadata"].cell(row=4, column=2).value, 2)
        finally:
            for name in list(sys.modules):
                if name == 'openpyxl' or name.startswith('openpyxl.'):
                    del sys.modules[name]
            sys.modules.update(stubbed)

    def test_save_output_ignores_direct_snapshot_that_lacks_diagnostics(self):
        """A direct-mode mid-fetch save has no fetch_strategy/direct diagnostics
        yet; letting it replace the persisted state would infer year mode and
        report complete for a paper still missing citations."""
        state = {
            "title": "Interrupted Paper",
            "fetch_strategy": "direct",
            "num_citations_on_scholar": 45,
            "complete_fetch_attempt": True,
            "direct_fetch_diagnostics": {
                "scholar_total": 45, "cached_total": 30, "seen_total": 30,
                "dedup_count": 0, "termination_reason": "in_progress",
            },
        }
        stale_partial = {
            "title": "Interrupted Paper",
            "num_citations_on_scholar": 45,
            "complete_fetch_attempt": False,
            "fetch_strategy": None,
            "direct_fetch_diagnostics": None,
            "year_fetch_diagnostics": {
                "histogram_total": 30, "seen_total": 30, "scholar_total": 45,
            },
            "citations": [],
        }
        pub = dict(self._INTERRUPTED_PUB, num_citations=45)
        paper, _ = self._save_and_load(
            state, with_snapshot=True, mid_snapshot=stale_partial, pub=pub)
        self.assertFalse(paper["fetch_complete"])
        self.assertEqual(paper["_fetch_state"]["fetch_strategy"], "direct")
        self.assertEqual(
            paper["_fetch_state"]["direct_fetch_diagnostics"]["seen_total"], 30)

    def test_interrupted_paper_is_partial_and_plannable_on_next_run(self):
        """Round trip for the reported bug: after an interrupted run the reloaded
        state must be 'partial' and show a real gap in the fetch plan, so the
        un-fetched citations are resumed instead of silently dropped."""
        from code.output_state import load_paper_states, to_paper_fetch_state
        _, _ = self._save_and_load(self._stale_complete_state(), with_snapshot=True)

        next_run_states = load_paper_states(self.fetcher.out_json)
        self.fetcher._paper_states = next_run_states

        pub = dict(self._INTERRUPTED_PUB)
        self.assertEqual(self.fetcher._citation_status(pub), "partial")

        state = next_run_states[pub["title"]]
        pst = to_paper_fetch_state(state)
        cached = pst.to_dict()
        cached["citations"] = list(state.citations)
        needed, seen, target, short_years = self.fetcher._plan_fetch_gap(
            "year", cached, pub["num_citations"], len(state.citations))
        self.assertEqual(needed, 1)
        self.assertEqual(short_years, "2026(4/5)")
        self.assertEqual([c["title"] for c in state.citations][:1], ["New-2026-0"])


if __name__ == '__main__':
    unittest.main()


class DedupKeepsSeenTotalTests(FetcherTestCase):
    """去重只改变"行怎么分组"，不改变"已经抓了多少"（用户 2026-10-15）。

    完整性判据只看 `seen_total`（对比 Scholar 的 `scholar_total` / `histogram_total`），
    所以输出去重**绝不能**把 `seen_total` 拉低 —— 否则"已经抓完"的论文下一轮会被当成
    没抓完而整篇重抓。同时要保持 `cached_total + dedup_count == seen_total`，否则统计口径
    自相矛盾（直接模式的 `dedup_count` 不会被读侧重算）。
    """

    PUB = {"no": 1, "title": "Some Paper", "num_citations": 2, "url": "https://x/1"}

    def _save_with_duplicates(self):
        row = {"title": "Citing Paper", "authors": "A", "venue": "CV",
               "year": "2024", "url": "https://example.com/cite", "cites_id": "12345"}
        fetch_state = {
            "fetch_strategy": "direct",
            "num_citations_on_scholar": 2,
            "direct_fetch_diagnostics": {
                "scholar_total": 2, "cached_total": 2, "seen_total": 2,
                "dedup_count": 0, "termination_reason": "iterator_exhausted",
            },
        }
        result = {"pub": dict(self.PUB),
                  "citations": [dict(row), dict(row)]}       # 同一 cites_id → 去重掉 1 行
        # `_build_entry` 的 fetch state 来自**运行时状态**（不是入参里的 `_fetch_state`）：
        # 直接种进 `_paper_states`（`to_paper_fetch_state` 接受 PaperFetchState 实例 ✓）。
        self.fetcher._paper_states = {self.PUB["title"]: PaperFetchState.from_dict(fetch_state)}
        with tempfile.TemporaryDirectory() as tmpdir:
            self.fetcher.profile_json = os.path.join(tmpdir, "author_test-author_profile.json")
            self.fetcher.out_json = os.path.join(tmpdir, "author_test-author_paper_citations.json")
            self.fetcher.out_xlsx = os.path.join(tmpdir, "author_test-author_paper_citations.xlsx")
            self.fetcher._run_start_time = 0
            with open(self.fetcher.profile_json, "w", encoding="utf-8") as handle:
                json.dump({"publications": [dict(self.PUB)]}, handle)
            self.fetcher._save_output([result])
            with open(self.fetcher.out_json, encoding="utf-8") as handle:
                return json.load(handle)["papers"][0]

    def test_dedup_moves_rows_into_dedup_count(self):
        entry = self._save_with_duplicates()
        self.assertEqual(len(entry["citations"]), 1, "重复行确实被收敛掉")
        diag = entry["_fetch_state"]["direct_fetch_diagnostics"]
        self.assertEqual(diag["seen_total"], 2, "已抓数量不许变")
        self.assertEqual(diag["cached_total"], 1)
        self.assertEqual(diag["dedup_count"], 1)
        self.assertEqual(diag["cached_total"] + diag["dedup_count"], diag["seen_total"])

    def test_paper_stays_complete_and_is_not_refetched(self):
        entry = self._save_with_duplicates()
        self.assertTrue(entry["fetch_complete"], "去重不许让论文变成未抓完")
        state = PaperFetchState.from_dict(entry["_fetch_state"])
        self.assertTrue(state.is_complete(2))
        self.assertFalse(state.need_fetch(2), "下一轮不该重抓")
